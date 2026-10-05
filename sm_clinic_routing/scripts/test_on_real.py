"""
test_on_real.py — Прогон реальных протоколов СМ-Клиники через систему.

Задачи:
  1. Читает все .docx из examples/protocols_real/
  2. Извлекает текст через python-docx
  3. Прогоняет через пайплайн (preprocessor → extractor → router)
  4. Считает метрики по органам
  5. Выводит таблицу + сохраняет JSON-отчёт

Запуск:
  python scripts/test_on_real.py
"""

from __future__ import annotations

import json
import logging
import sys
import time
from collections import defaultdict
from dataclasses import asdict, dataclass, field
from pathlib import Path

# Добавляем src/ в путь
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))

from config_loader import ConfigLoader
from extractor import Extractor
from router import Router

logging.basicConfig(
    level=logging.WARNING,  # только предупреждения
    format="%(asctime)s [%(levelname)s] %(name)s: %(message)s",
)
log = logging.getLogger("test_on_real")


# ============================================================================
#                        ИЗВЛЕЧЕНИЕ ТЕКСТА
# ============================================================================

# Папка для кэша
CACHE_DIR = Path(__file__).resolve().parent.parent / "examples" / ".cache"


def extract_text_from_docx(path: Path, use_cache: bool = True) -> str:
    """
    Извлекает текст из .docx через python-docx.
    Склеивает абзацы и содержимое таблиц.

    Если use_cache=True и кэш существует — читает .txt вместо .docx.
    Кэш инвалидируется по mtime .docx.
    """
    # 1. Проверяем кэш
    if use_cache:
        cache_file = CACHE_DIR / f"{path.stem}.txt"
        meta_file = CACHE_DIR / f"{path.stem}.meta"

        if cache_file.exists() and meta_file.exists():
            try:
                cached_mtime = float(meta_file.read_text(encoding="utf-8").strip())
                actual_mtime = path.stat().st_mtime
                if abs(cached_mtime - actual_mtime) < 0.001:
                    return cache_file.read_text(encoding="utf-8")
            except (OSError, ValueError):
                log.debug(f"Кэш битый для {path.name} — парсим заново")

    # 2. Парсим .docx
    try:
        from docx import Document
    except ImportError:
        return ""

    try:
        doc = Document(str(path))
    except (OSError, ValueError, KeyError) as e:
        log.warning(f"Не удалось открыть {path.name}: {e}")
        return ""

    parts: list[str] = []

    # Абзацы
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            parts.append(text)

    # Таблицы
    for table in doc.tables:
        for row in table.rows:
            cells_text = []
            for cell in row.cells:
                cell_text = cell.text.strip()
                if cell_text:
                    cells_text.append(cell_text)
            if cells_text:
                parts.append("; ".join(cells_text))

    result = "\n".join(parts)

    # 3. Сохраняем в кэш
    if use_cache and result:
        try:
            CACHE_DIR.mkdir(parents=True, exist_ok=True)
            cache_file = CACHE_DIR / f"{path.stem}.txt"
            meta_file = CACHE_DIR / f"{path.stem}.meta"
            cache_file.write_text(result, encoding="utf-8")
            meta_file.write_text(str(path.stat().st_mtime), encoding="utf-8")
        except (OSError, ValueError) as e:
            log.warning(f"Не удалось сохранить кэш для {path.name}: {e}")

    return result


# ============================================================================
#                        DATACLASSES ДЛЯ МЕТРИК
# ============================================================================


@dataclass
class ProtocolResult:
    """Результат обработки одного протокола."""

    file: str
    organ_folder: str
    text_length: int = 0
    sentences_total: int = 0

    # Результат
    findings_total: int = 0
    findings_positive: int = 0
    findings_negative: int = 0

    # Маршрут
    status: str = "unknown"  # assigned / no_action
    specialist: list[str] = field(default_factory=list)
    urgency: str = ""
    primary_finding_id: str = ""
    primary_finding_synonym: str = ""
    primary_organ: str = ""

    # Скорость
    duration_ms: float = 0.0

    # Ошибка
    error: str = ""


@dataclass
class Metrics:
    """Общие метрики по всем протоколам."""

    total_protocols: int = 0
    total_assigned: int = 0  # создан маршрут
    total_no_action: int = 0  # норма
    total_errors: int = 0

    by_organ: dict[str, dict] = field(default_factory=dict)
    total_duration_ms: float = 0.0
    avg_duration_ms: float = 0.0

    results: list[ProtocolResult] = field(default_factory=list)


# ============================================================================
#                        ОСНОВНАЯ ЛОГИКА
# ============================================================================


def process_one(
    docx_path: Path,
    organ_folder: str,
    ext: Extractor,
    rt: Router,
) -> ProtocolResult:
    """Обрабатывает один протокол, возвращает результат."""
    result = ProtocolResult(
        file=docx_path.name,
        organ_folder=organ_folder,
    )

    # 1. Извлекаем текст
    text = extract_text_from_docx(docx_path)
    result.text_length = len(text)

    if not text:
        result.error = "Не удалось извлечь текст"
        return result

    # 2. Прогоняем через пайплайн
    start = time.perf_counter()

    try:
        extraction = ext.extract(text)
        route = rt.build(
            extraction,
            study_type=f"УЗИ ({organ_folder})",
            study_date="",
        )
    except (ValueError, KeyError, AttributeError, TypeError) as e:
        result.error = f"{type(e).__name__}: {e}"
        return result

    duration_ms = (time.perf_counter() - start) * 1000
    result.duration_ms = round(duration_ms, 2)

    # 3. Метрики находок
    result.sentences_total = extraction.sentences_total
    result.findings_total = len(extraction.findings)
    result.findings_positive = len(extraction.positive())
    result.findings_negative = len(extraction.negative())

    # 4. Маршрут
    result.status = route.status
    result.specialist = list(route.specialist)
    result.urgency = route.urgency

    if route.primary_finding:
        result.primary_finding_id = route.primary_finding.id
        result.primary_finding_synonym = route.primary_finding.matched_synonym
        result.primary_organ = route.primary_finding.organ_code

    return result


# ============================================================================
#                        СБОР МЕТРИК
# ============================================================================


def collect_metrics(results: list[ProtocolResult]) -> Metrics:
    """Считает метрики по всем результатам."""
    metrics = Metrics()
    metrics.results = results
    metrics.total_protocols = len(results)

    by_organ: dict[str, dict] = defaultdict(
        lambda: {
            "total": 0,
            "assigned": 0,
            "no_action": 0,
            "errors": 0,
            "findings_total": 0,
            "duration_ms": 0.0,
        }
    )

    for r in results:
        organ = r.organ_folder
        by_organ[organ]["total"] += 1

        if r.error:
            metrics.total_errors += 1
            by_organ[organ]["errors"] += 1
            continue

        if r.status == "assigned":
            metrics.total_assigned += 1
            by_organ[organ]["assigned"] += 1
        elif r.status == "no_action":
            metrics.total_no_action += 1
            by_organ[organ]["no_action"] += 1

        by_organ[organ]["findings_total"] += r.findings_total
        by_organ[organ]["duration_ms"] += r.duration_ms
        metrics.total_duration_ms += r.duration_ms

    if metrics.total_protocols:
        metrics.avg_duration_ms = round(
            metrics.total_duration_ms / metrics.total_protocols, 2
        )

    # Округляем
    for organ, data in by_organ.items():
        if data["total"]:
            data["duration_ms"] = round(data["duration_ms"], 2)

    metrics.by_organ = dict(by_organ)
    return metrics


def print_report(metrics: Metrics) -> None:
    """Красивый вывод таблицы в консоль."""
    print()
    print("=" * 80)
    print("  МЕТРИКИ ОБРАБОТКИ РЕАЛЬНЫХ ПРОТОКОЛОВ СМ-КЛИНИКИ")
    print("=" * 80)
    print()

    # Таблица по органам
    header = f"{'Орган':<35} {'Всего':>7} {'Assigned':>9} {'NoAction':>9} {'Ошибок':>7}"
    print(header)
    print("-" * 80)

    for organ, data in sorted(metrics.by_organ.items()):
        organ_short = organ.replace("протоколы ", "").strip()
        print(
            f"{organ_short:<35} "
            f"{data['total']:>7} "
            f"{data['assigned']:>9} "
            f"{data['no_action']:>9} "
            f"{data['errors']:>7}"
        )

    print("-" * 80)
    print(
        f"{'ИТОГО':<35} "
        f"{metrics.total_protocols:>7} "
        f"{metrics.total_assigned:>9} "
        f"{metrics.total_no_action:>9} "
        f"{metrics.total_errors:>7}"
    )
    print()

    # Сводка
    if metrics.total_protocols:
        rate = metrics.total_assigned / metrics.total_protocols * 100
        print(f"📊 Всего протоколов:            {metrics.total_protocols}")
        print(
            f"✅ Маршрут создан:               {metrics.total_assigned} ({rate:.1f}%)"
        )
        print(f"🔵 Без находок (no_action):      {metrics.total_no_action}")
        print(f"❌ Ошибок обработки:             {metrics.total_errors}")
        print()
        print(f"⏱️  Всего времени:                {metrics.total_duration_ms:.0f} мс")
        print(f"⏱️  Среднее на протокол:          {metrics.avg_duration_ms:.1f} мс")
        print(
            f"⏱️  Скорость:                    {1000 / metrics.avg_duration_ms:.0f} протоколов/сек"
            if metrics.avg_duration_ms > 0
            else ""
        )
    print()
    print("=" * 80)


def save_json_report(metrics: Metrics, output_path: Path) -> None:
    """Сохраняет JSON-отчёт."""
    report = {
        "summary": {
            "total_protocols": metrics.total_protocols,
            "total_assigned": metrics.total_assigned,
            "total_no_action": metrics.total_no_action,
            "total_errors": metrics.total_errors,
            "avg_duration_ms": metrics.avg_duration_ms,
            "total_duration_ms": round(metrics.total_duration_ms, 2),
        },
        "by_organ": metrics.by_organ,
        "details": [asdict(r) for r in metrics.results],
    }

    output_path.parent.mkdir(parents=True, exist_ok=True)
    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(report, f, ensure_ascii=False, indent=2)

    print(f"💾 JSON-отчёт сохранён: {output_path}")


# ============================================================================
#                        ГЛАВНАЯ ФУНКЦИЯ
# ============================================================================


def main() -> None:
    # 1. Пути
    root = Path(__file__).resolve().parent.parent
    protocols_dir = root / "examples" / "protocols_real"
    output_dir = root / "output"
    output_dir.mkdir(parents=True, exist_ok=True)

    if not protocols_dir.exists():
        print(f"❌ Папка не найдена: {protocols_dir}")
        sys.exit(1)

    # 2. Инициализация
    print("🔧 Загрузка конфигов...")
    cfg = ConfigLoader(str(root / "config")).load()
    print(f"✅ Загружено: {len(cfg.organs)} органов")

    ext = Extractor(cfg)
    rt = Router(cfg)

    # 3. Собираем файлы
    docx_files: list[tuple[Path, str]] = []
    for organ_dir in sorted(protocols_dir.iterdir()):
        if not organ_dir.is_dir():
            continue
        organ_name = organ_dir.name
        for f in sorted(organ_dir.glob("*.docx")):
            docx_files.append((f, organ_name))

    total = len(docx_files)
    if not total:
        print(f"❌ Не найдено .docx файлов в {protocols_dir}")
        sys.exit(1)

    print(f"📂 Найдено протоколов: {total}")
    print()

    # 4. Обрабатываем
    results: list[ProtocolResult] = []
    for i, (docx_path, organ_folder) in enumerate(docx_files, 1):
        # Прогресс каждые 10 протоколов
        if i % 10 == 0 or i == 1 or i == total:
            print(f"  [{i}/{total}] {organ_folder} / {docx_path.name}")

        result = process_one(docx_path, organ_folder, ext, rt)
        results.append(result)

        if result.error:
            print(f"     ⚠️  {result.error}")

    print()

    # 5. Считаем метрики
    metrics = collect_metrics(results)

    # 6. Выводим отчёт
    print_report(metrics)

    # 7. Сохраняем JSON
    save_json_report(metrics, output_dir / "metrics_report.json")

    print()
    print("✅ ГОТОВО!")


if __name__ == "__main__":
    main()
