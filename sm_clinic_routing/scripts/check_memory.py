"""
check_memory.py — Проверка утечек памяти через tracemalloc.

Аналог valgrind для Python.
Запускает пайплайн на 10 протоколах и считает пиковое потребление.

Запуск:
    python scripts/check_memory.py
"""

import sys
import tracemalloc
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))

from config_loader import ConfigLoader
from extractor import Extractor
from router import Router


def extract_text(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                t = cell.text.strip()
                if t:
                    parts.append(t)
    return "\n".join(parts)


def main():
    print("🧠 Проверка памяти через tracemalloc")
    print("=" * 60)

    cfg = ConfigLoader(str(ROOT / "config")).load()
    ext = Extractor(cfg)
    rt = Router(cfg)

    # Берём 10 протоколов
    files = list((ROOT / "examples" / "protocols_real").rglob("*.docx"))[:10]
    print(f"📂 Протоколов: {len(files)}")

    # Запуск
    tracemalloc.start()

    for path in files:
        try:
            text = extract_text(path)
            extraction = ext.extract(text)
            rt.build(extraction, study_type="УЗИ", study_date="")
        except (OSError, ValueError, KeyError) as e:
            print(f"⚠️  {path.name}: {e}")

    current, peak = tracemalloc.get_traced_memory()

    print()
    print("=" * 60)
    print(f"📊 Текущая память:  {current / 1024:.1f} КБ")
    print(f"📊 Пиковая память:  {peak / 1024 / 1024:.2f} МБ")
    print("=" * 60)

    # Топ-5 по памяти — ДО остановки tracemalloc
    print()
    print("Топ-5 функций по памяти:")
    snapshot = tracemalloc.take_snapshot()
    top_stats = snapshot.statistics("lineno")
    for stat in top_stats[:5]:
        print(f"  {stat}")

    # Останавливаем ПОСЛЕ snapshot
    tracemalloc.stop()


if __name__ == "__main__":
    main()
