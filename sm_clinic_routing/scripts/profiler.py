"""
profile.py — Профилирование узких мест в обработке протокола.

Запуск:
    python scripts/profile.py
"""

import cProfile
import pstats
import sys
import time
from io import StringIO
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))

from config_loader import ConfigLoader
from extractor import Extractor
from router import Router

# ============================================================================
#  Извлечение текста (как в test_on_real.py)
# ============================================================================


def extract_text_from_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    parts = []
    for para in doc.paragraphs:
        t = para.text.strip()
        if t:
            parts.append(t)
    for table in doc.tables:
        for row in table.rows:
            cells = []
            for cell in row.cells:
                ct = cell.text.strip()
                if ct:
                    cells.append(ct)
            if cells:
                parts.append("; ".join(cells))
    return "\n".join(parts)


# ============================================================================
#  MAIN
# ============================================================================


def main():
    # Берём один протокол ЩЖ (там точно есть находки)
    protocol_path = None
    for p in Path("examples/protocols_real/протоколы щитовидная железа").glob("*.docx"):
        protocol_path = p
        break

    if not protocol_path:
        print("❌ Не найден протокол ЩЖ")
        return

    print(f"📄 Протокол: {protocol_path.name}")

    # 1. Парсинг .docx
    t0 = time.perf_counter()
    text = extract_text_from_docx(protocol_path)
    t1 = time.perf_counter()
    print(f"  1. Парсинг .docx:   {(t1-t0)*1000:6.1f} мс (символов: {len(text)})")

    # 2. Загрузка конфигов
    t0 = time.perf_counter()
    cfg = ConfigLoader("config").load()
    t1 = time.perf_counter()
    print(f"  2. ConfigLoader:    {(t1-t0)*1000:6.1f} мс (органов: {len(cfg.organs)})")

    ext = Extractor(cfg)
    rt = Router(cfg)

    # 3. Preprocessor
    from preprocessor import Preprocessor

    pre = Preprocessor(cfg)

    t0 = time.perf_counter()
    pre_result = pre.process(text)
    t1 = time.perf_counter()
    print(
        f"  3. Preprocessor:    {(t1-t0)*1000:6.1f} мс (предложений: {len(pre_result.sentences)})"
    )

    # 4. Extractor
    t0 = time.perf_counter()
    extraction = ext.extract(text)
    t1 = time.perf_counter()
    print(
        f"  4. Extractor:       {(t1-t0)*1000:6.1f} мс (находок: {len(extraction.findings)})"
    )

    # 5. Router
    t0 = time.perf_counter()
    route = rt.build(extraction, study_type="УЗИ ЩЖ", study_date="")
    t1 = time.perf_counter()
    print(f"  5. Router:          {(t1-t0)*1000:6.1f} мс (статус: {route.status})")

    print()

    # 6. Полный профиль Extractor (cProfile)
    print("=" * 60)
    print("ПРОФИЛЬ Extractor.extract() — топ-15 функций по времени")
    print("=" * 60)

    profiler = cProfile.Profile()
    profiler.enable()
    ext.extract(text)
    profiler.disable()

    s = StringIO()
    ps = pstats.Stats(profiler, stream=s).sort_stats("cumulative")
    ps.print_stats(15)
    print(s.getvalue())


if __name__ == "__main__":
    main()
